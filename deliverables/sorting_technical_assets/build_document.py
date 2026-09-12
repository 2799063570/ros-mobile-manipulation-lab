from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
ASSETS = Path(__file__).resolve().parent
OUT = ROOT/'deliverables/AUBO_i5智能分拣与视觉伺服系统_技术文档.docx'
FONT = 'C:/Windows/Fonts/msyh.ttc'

def diagram(name, boxes, arrows, size=(1600,850), notes=()):
    im=Image.new('RGB',size,'white'); d=ImageDraw.Draw(im)
    font=ImageFont.truetype(FONT,27); small=ImageFont.truetype(FONT,23)
    for a,b,*label in arrows:
        d.line([a,b],fill='#557084',width=4)
        import math
        ang=math.atan2(b[1]-a[1],b[0]-a[0]); r=13
        d.polygon([b,(b[0]-r*math.cos(ang-.5),b[1]-r*math.sin(ang-.5)),(b[0]-r*math.cos(ang+.5),b[1]-r*math.sin(ang+.5))],fill='#557084')
        if label: d.text(((a[0]+b[0])/2+12,(a[1]+b[1])/2-30),label[0],font=small,fill='#445566')
    for rect,txt,*opts in boxes:
        d.rounded_rectangle(rect,radius=15,fill=opts[0] if opts else '#EDF3F7',outline='#7590A4',width=2)
        x0,y0,x1,y1=rect
        bb=d.multiline_textbbox((0,0),txt,font=font,spacing=10,align='center')
        d.multiline_text(((x0+x1-bb[2])/2,(y0+y1-(bb[3]-bb[1]))/2-bb[1]),txt,font=font,fill='#172B3A',spacing=10,align='center')
    for xy,txt in notes:d.text(xy,txt,font=small,fill='#445566')
    im.save(ASSETS/name)

diagram('architecture.png',[
 ((400,20,1200,110),'RealSense D435i  RGB 与对齐深度  默认 30 Hz'),
 ((400,160,1200,260),'YOLO-OBB  默认限频 10 Hz\n深度配对  内参反投影  手眼外参变换'),
 ((45,355,735,450),'分拣链路  DetectedObjectArray\n类别  中心位置  抓取角度  开口宽度'),
 ((865,355,1555,450),'伺服链路  PoseStamped\n目标位置与采集时间戳'),
 ((45,505,735,600),'任务状态机 → MoveIt → 轨迹 Action\n按任务触发  RRT-Connect 与笛卡尔路径'),
 ((865,505,1555,600),'VisualServo  100 Hz\n混合模式  MoveIt规划服务＋近距PBVS'),
 ((45,660,735,765),'RobotHW 主循环 250 Hz\n轨迹缓存 → SDK 独立供点线程'),
 ((865,660,1555,765),'有界队列 → DirectSdkBackend\n200 Hz 点间隔模型  批量供点'),
], [((800,110),(800,160)),((650,260),(390,355)),((950,260),(1210,355)),((390,450),(390,505)),((1210,450),(1210,505)),((390,600),(390,660)),((1210,600),(1210,660))],size=(1600,845),notes=[((80,797),'真机两条链路互斥使用控制柜；相同 SDK 不代表相同控制进程或同一条队列。')])

diagram('perception.png',[
 ((35,35,485,150),'RGB 图像\nOBB 中心 u v  宽高  角度'),
 ((560,35,1010,150),'对齐深度缓存\n按 RGB 原始时间戳匹配'),
 ((1085,35,1535,150),'CameraInfo 与手眼标定\n内参 K  外参 T_base_camera'),
 ((300,245,1300,355),'有效深度筛选与中值 → 三维反投影 → 基座坐标'),
 ((80,455,750,570),'位置分支\n顶面高度 → 物体中心 → TCP 抓取偏移'),
 ((850,455,1520,570),'方向分支\n旋转框中线投影 → 短边方向 → 夹爪 yaw'),
 ((300,665,1300,765),'几何有效性  工作空间  多帧稳定性检查 → 抓取候选'),
], [((260,150),(500,245)),((785,150),(800,245)),((1310,150),(1100,245)),((600,355),(415,455)),((1000,355),(1185,455)),((415,570),(550,665)),((1185,570),(1050,665))],size=(1600,795))

diagram('servo.png',[
 ((35,35,485,135),'目标位姿\n新鲜度检查  TF 变换'),
 ((565,35,1015,135),'关节反馈 50 Hz\nKDL 正运动学与雅可比'),
 ((1095,35,1555,135),'设定值与约束\n目标偏移  姿态  速度上限'),
 ((350,220,1250,320),'100 Hz 控制定时回调\n位置软死区  姿态误差 → 笛卡尔速度 → DLS'),
 ((350,405,1250,505),'限关节速度与加速度  关节位置边界\n每周期约 2 个子步  5 ms 位置积分'),
 ((350,590,1250,680),'CommandQueue  标准80点  混合8点\n生产与输出解耦  满时淘汰最旧待执行点'),
 ((350,760,1250,860),'SDK 独立输出线程\n诊断与缓冲查询 → 二次限幅 → 批量写入'),
], [((260,135),(500,220)),((790,135),(800,220)),((1325,135),(1100,220)),((800,320),(800,405)),((800,505),(800,590)),((800,680),(800,760))],size=(1600,905))

diagram('hybrid.png',[
 ((340,15,1260,105),'初始化与反馈保持  等待新鲜目标'),
 ((340,150,1260,240),'距离较远 → PLANNING  10 Hz检查待规划请求'),
 ((340,285,1260,390),'MoveIt服务规划 → 校验轨迹 → APPROACH\n100 Hz路径采样  共用队列与SDK输出'),
 ((340,440,1260,535),'达到接近点并稳定 → 近距PBVS 100 Hz\n进入阈值0.10 m  退出阈值0.16 m'),
 ((340,585,1260,685),'TRACKING → ALIGNED  对准位保持\n目标漂移或超时 → 停止旧路径并重新评估'),
 ((340,735,1260,825),'后续任务衔接  接近与夹取  搬运  放置\n完整抓放仍由任务层组织'),
], [((800,105),(800,150)),((800,240),(800,285)),((800,390),(800,440)),((800,535),(800,585)),((800,685),(800,735))],size=(1600,880),notes=[((20,850),'目标超时、规划失败或 SDK 故障时，转入保持或失败处理，不继续抓取。')])

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.7); sec.bottom_margin=Cm(1.6); sec.left_margin=Cm(1.8); sec.right_margin=Cm(1.8)
sec.footer_distance=Cm(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    s=doc.styles[name]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    s.font.size=Pt(10.5)
    s.paragraph_format.space_after=Pt(6)
    s.paragraph_format.line_spacing=1.18
doc.styles['Title'].font.size=Pt(23); doc.styles['Title'].font.bold=True
doc.styles['Heading 1'].font.size=Pt(17); doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 2'].font.size=Pt(12); doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
doc.styles['Caption'].font.size=Pt(9); doc.styles['Caption'].paragraph_format.space_after=Pt(9)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('AUBO i5   |   '); r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

def p(t,style=None):return doc.add_paragraph(t,style)
def h(t):doc.add_heading(t,2)
def page(t):doc.add_page_break();doc.add_heading(t,1)
def pic(name,w,cap):
    q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(3)
    q.add_run().add_picture(str(ASSETS/name),width=Cm(w));p(cap,'Caption')
def table(headers,rows,widths):
    if headers[0]=='编号':
        rows=sorted(rows,key=lambda row:int(row[0][1:]))
    if headers[0]=='执行单元':
        table(['执行单元及职责']+headers[1:],rows[:11],widths)
        h('线程与频率的区别')
        p('感知与任务节点按消息或请求工作，不能为每个函数指定固定Hz。异步线程池允许状态读取、停止请求和任务回调并发处理；独立SDK线程负责网络诊断与供点。规划耗时与动作执行时间由任务复杂度决定。')
        page('五 视觉伺服线程与时序')
        p('视觉伺服在一个节点中组合控制定时器、反馈定时器、输出后端和异步回调线程池。混合模式另外启用规划检查定时器，规划服务调用期间释放控制锁，避免阻塞其余控制回调。')
        table(['执行单元及职责']+headers[1:],rows[11:],widths)
        h('100 Hz控制与200 Hz位置点')
        p('一个名义10 ms控制周期内，先以目标误差求一次关节速度，再分两个约5 ms子步完成约束与位置积分。Gazebo输出按200 Hz消费；SDK线程按缓冲需求批量写点，每轮末尾休眠4 ms，并不意味着每轮只执行一个位置点。')
        p('标准模式的80点软件队列，在200点/秒的名义节拍下对应最多约0.4 s的点数规模；混合模式缩短为8点，约0.04 s。这只是队列容量对应的时间，不是实际延迟；总延迟还取决于控制柜缓存与网络。')
        return
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for i,v in enumerate(headers):t.rows[0].cells[i].text=v
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):cells[i].text=str(v)
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if ri==0:trPr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DBE7EF' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF'));tcPr.append(sh)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcPr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            tcPr.append(margins)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(2);par.paragraph_format.space_before=Pt(2);par.paragraph_format.line_spacing=1.05
                for run in par.runs:run.font.size=Pt(9);run.bold=ri==0
    p('') .paragraph_format.space_after=Pt(0)
    return t

p('AUBO i5机械臂\n智能分拣与视觉伺服系统','Title')
p('招聘展示技术文档   2025.04 – 2026.08','Subtitle')
p('我们面向目标位置和摆放方向不确定的分拣场景，构建从旋转目标检测、RGB-D几何定位到机械臂规划执行的系统，并开发基于位姿误差的视觉伺服模块。系统以眼在手外的RealSense D435i为主要展示方案，通过“长距离运动规划＋近距离视觉伺服”的设计改善单次视觉定位后缺少在线修正的问题。')
p('本文重点展示感知与抓取几何、MoveIt执行链、SDK硬件接口、视觉闭环算法以及线程协作机制，便于招聘交流时结合代码和演示说明工程实现。')
table(['项目指标','结果','指标含义'],[
 ('旋转目标检测','mAP@0.5  77.2%','检测精度，不等同于抓取成功率'),
 ('目标位置与姿态估计误差','5.8 mm / 3.3°','定位与方向估计指标'),
 ('平均预测耗时','5.4 ms','模型预测阶段耗时'),
 ('单次运动规划耗时','< 100 ms','实验场景下的规划结果'),
 ('综合抓取成功率','92.6%','完整抓取任务的结果指标'),
],[4.8,4,8.6])
p('指标为项目实验记录值；模型耗时、整条感知链延迟、控制周期和任务成功率采用不同统计口径，应分别展示。详细测试口径见第八节。')
pic('image2.png',12.2,'图1  原项目文档中的旋转目标检测画面  展示瓶体 罐体与盒体的类别和方向')
p('技术栈  ROS 1 · C++ / Python · YOLO-OBB · RGB-D · TF · MoveIt / OMPL · ros_control · AUBO SDK · KDL / Eigen')

page('一 系统架构与模块职责')
p('系统按感知、任务规划和机械臂执行分层。检测器负责回答“是什么、在哪里、朝向如何”；几何层将像素结果变成机器人可用的抓取参数；任务层组织抓取顺序；执行层负责轨迹或伺服位置指令与控制柜的通信。')
pic('architecture.png',17.2,'图2  系统运行架构  频率为当前配置或源码名义值  独立RobotHW与伺服SDK输出互斥')
table(['模块','仓库位置','核心职责'],[
 ('感知与几何','aubo/aubo_perception','模型接入、深度配对、坐标变换、抓取角度与宽度'),
 ('分拣任务','aubo/aubo_sorting_core','目标稳定、抓放状态机、失败返回与动作执行'),
 ('规划与模型','aubo/aubo_moveit_config\naubo/aubo_description','URDF、关节与碰撞约束、OMPL配置'),
 ('控制与驱动','aubo/aubo_ros_control','RobotHW、视觉伺服、队列、SDK通信'),
 ('场景配置','aubo/aubo_sorting','固定机械臂工位、启动文件与抓取参数'),
],[3,5.4,9])
h('实现边界')
p('常规分拣通过MoveIt Action与RobotHW执行；混合控制则在同一个VisualServo节点中请求MoveIt规划服务，将远距离轨迹和近距离PBVS共同送入伺服队列与同一输出后端，避免两套SDK同时控制。完整夹取、搬运和放置仍由分拣任务层组织。[S9]')

page('二 从旋转检测到机械臂抓取位姿')
p('YOLO适配节点加载外部模型工程中的GC-yolo.pt，解析OBB中心、宽高、方向角、类别与置信度，发布统一检测消息。模型结构改进属于独立训练工程；ROS侧的重点是把检测输出可靠地接入实时几何与抓取流程。[S1]')
pic('perception.png',16.5,'图3  抓取几何流程  图像方向必须经过相机投影与坐标变换')
h('深度与时间对齐')
p('检测消息保留RGB采集时间戳，几何节点按该时间从深度缓存中匹配图像。默认缓存90帧、允许时间差0.08 s；旋转框中心区域筛除无效及越界深度后取中值，有效样本少于5个则拒绝该目标。这样可减小孔洞、背景和推理延迟引入的定位偏差。[S2]')
h('反投影与手眼变换')
p('对中心像素(u, v)与深度Z，计算X = (u − cx)Z / fx、Y = (v − cy)Z / fy，再由p_base = R_base_camera p_camera + t_base_camera转换至机械臂基座系。内参来自CameraInfo；眼在手外的外参由标定结果发布，安装位置变化后必须重新标定。')
h('位置和方向的工程处理')
p('深度读数对应物体可见顶面。代码先校验顶面与桌高加物高是否一致，再减去半个物高得到物体中心，最后叠加TCP抓取偏移。另有table模式，用相机射线与已知水平面求交；该模式依赖桌高和物高先验，不能与深度定位混为一谈。')
p('方向计算将旋转框两条中线投影到基座水平面，选取米制短边作为夹爪闭合方向，再叠加夹爪轴标定偏置。该方法表达桌面抓取位置和偏航角，并非任意物体完整6D姿态重建。抓取宽度控制还需要夹爪开口标定，默认不启用。')

page('三 MoveIt规划与SDK执行链')
h('抓放任务如何组织')
p('任务先完成初始化与观察位运动，再积累目标检测样本并检查稳定性。当前目标缓存按类别组织，每类维护一个候选；它适合当前分拣工位，不应描述为同类多目标身份跟踪。通过检查的目标进入预抓取、接近、闭合、抬升、搬运、放置和退离流程。[S3]')
table(['阶段','执行方式','关键检查'],[
 ('预抓取与搬运','MoveIt位姿规划','关节限制、场景碰撞、规划与执行返回值'),
 ('接近与放置','笛卡尔路径及时间参数化','路径完成比例；必要时按代码分支重试或回退'),
 ('抬升','抬升路径与恢复逻辑','不把未完成的部分抬升当成达到安全高度'),
 ('夹取与释放','夹爪FollowJointTrajectory','动作超时、位置容差与抓取几何合法性'),
],[3.7,5.1,8.6])
h('运动规划设置')
p('URDF与SRDF提供机器人运动链、碰撞模型和规划组；关节范围、速度与加速度缩放约束运动。OMPL配置中默认规划器为RRTConnect；长距离运动由其求解，局部直线动作通过computeCartesianPath实现。规划场景中的桌面和障碍物参与碰撞检查，点云与OctoMap能力可按场景接入。[S4]')
p('配置中的planning_time为12 s，表示允许的规划时间预算。“单次规划<100 ms”是实验结果，两者并不冲突，也不能用12 s配置或单次快速成功来证明所有工况都小于100 ms。')
h('ros_control硬件接口')
p('我们基于AUBO SDK封装AuboHardwareInterface，继承RobotHW并注册关节状态接口与位置命令接口。MoveIt将带时间参数的轨迹交给FollowJointTrajectory动作接口，JointTrajectoryController在controller_manager更新中生成各关节位置命令。[S5]')
p('主线程以默认250 Hz执行read → update → write。SDK状态定时器以50 Hz更新反馈缓存，read读取该缓存；write进行指令合法性与步长检查，并将需要执行的位置点推入单生产者单消费者队列。独立供点线程查询控制柜缓冲余量，批量取点、执行约束检查后下发。')
h('线程解耦的作用')
p('轨迹计算、关节状态读取和控制柜供点分开调度，降低网络查询阻塞控制器更新的影响。批量写入减少通信调用次数，但供点轮询频率不等于控制柜执行采样率。该路径的限幅常量CTRL_PERIOD_S为5 ms，与250 Hz主循环的4 ms周期不同，修改频率前应同时核对点间隔与控制柜实际消费节拍。')

page('四 视觉伺服算法与前向位置队列')
p('视觉伺服采用基于三维位姿误差的PBVS。眼在手外时，期望TCP位置等于目标基座坐标加抓取偏移；当前TCP位置和雅可比由关节反馈及KDL运动学链计算。控制器输出关节速度，再积分为SDK接受的关节位置点。[S6]')
pic('servo.png',13.8,'图4  伺服计算与SDK执行  控制回调100 Hz  名义位置点间隔5 ms')
h('控制律')
p('位置误差 e_p = p_target + p_offset − p_TCP。逐轴应用连续软死区后，v = K_p e_p；姿态误差取旋转矩阵对数，ω = K_R Log(R_desired R_TCPᵀ)。笛卡尔速度先限幅，再通过阻尼最小二乘得到 q̇ = Jᵀ(JJᵀ + λ²I)⁻¹[v; ω]。阻尼项减轻近奇异位形下逆解的数值放大。')
p('关节速度经过速度与加速度限幅，以 q_next = q_cmd + q̇ Δt积分并限制在关节位置边界内。100 Hz控制回调按实际间隔生成子步，默认每10 ms约生成两个5 ms位置点；少量反馈融合抑制积分位置与实际关节位置长期偏离。这里是运动学约束与位置积分，不是完整动力学或力矩控制。')
h('队列与姿态含义')
p('标准伺服队列容量为80点，混合模式配置缩短为8点；满时淘汰最旧待执行点；状态切换时清除失效命令，输出端再次限速和限加速度。眼在手外当前姿态控制跟踪配置的固定抓取姿态，YOLO伺服目标的位置接口没有直接承载OBB抓取yaw；分拣链路中的角度估计与伺服姿态控制应分别说明。')

page('五 线程 回调与执行频率')
p('下表列出源码和默认配置中的运行节拍，不代表实机测得的稳定频率。ROS定时器由回调线程池执行；每个定时器并不对应一个独立线程。SDK供点循环包含通信时间和休眠，因此“4 ms休眠”只能表示约250 Hz的轮询上限。[S5–S8]')
table(['执行单元','调度方式','默认频率或周期','作用'],[
 ('RealSense RGB与深度','驱动采集','30 Hz / 33.3 ms','真机启动沿用相机驱动默认值'),
 ('固定相机仿真','Gazebo传感器','20 Hz / 50 ms','workspace_camera的update_rate'),
 ('YOLO推理','图像回调＋限频锁','最高10 Hz / ≥100 ms','单调时钟限频；忙时跳过'),
 ('RGB-D几何计算','检测消息回调','随有效检测到达','不是独立固定频率循环'),
 ('目标稳定等待','任务线程内部轮询','10 Hz / 100 ms','waitForObject等等待逻辑'),
 ('任务初始化与抓放','独立工作线程','按任务触发','规划和动作等待，不固定Hz'),
 ('任务ROS回调','AsyncSpinner(4)','4个工作线程','线程数不代表4 Hz'),
 ('RobotHW主控制','主线程循环','250 Hz / 4 ms','read → update → write'),
 ('RobotHW状态读取','ROS定时回调','50 Hz / 20 ms','仅状态模式降低为1 Hz'),
 ('RobotHW SDK供点','独立线程','每轮末休眠4 ms','缓冲不足时批量补点'),
 ('硬件节点ROS回调','AsyncSpinner(2)','2个工作线程','状态定时器和服务回调'),
 ('视觉伺服控制','ROS定时回调','100 Hz / 10 ms','误差、逆解与子步积分'),
 ('视觉伺服状态读取','SDK模式定时回调','50 Hz / 20 ms','读取并发布关节反馈'),
 ('Gazebo伺服输出','ROS定时回调','200 Hz / 5 ms','消费队列并发布关节位置'),
 ('SDK伺服输出','独立线程','每轮末休眠4 ms','位置点按200 Hz模型限幅'),
 ('混合规划检查','ROS定时回调','10 Hz / 100 ms','有请求才调用MoveIt服务'),
 ('伺服ROS回调及监测','AsyncSpinner(3)＋主线程','3个回调线程；监测10 Hz','监测不是控制计算循环'),
],[3.7,3.6,4.1,6])
h('控制器附属发布频率')
p('controllers.yaml中JointStateController发布频率为50 Hz；轨迹控制器状态发布为25 Hz、Action监测为20 Hz。它们分别用于状态输出和动作状态监测，不改变RobotHW主循环的默认250 Hz。')
h('视觉频率与控制频率如何配合')
p('默认10 Hz视觉观测之间，100 Hz控制器可能使用同一帧新鲜目标多次计算。5.4 ms平均模型预测耗时不等于185 Hz整机视觉闭环；实际链路还包含采集、排队、深度配对、TF、控制回调和SDK缓冲延迟。')

page('六 混合控制流程与状态切换')
p('混合控制在同一个VisualServo节点中完成远距离规划与近距离PBVS切换。通过/plan_kinematic_path只请求规划结果，校验关节顺序、时间、限位及起终点后，以统一队列输出轨迹采样点；该路径不使用MoveIt轨迹Action执行。[S9]')
pic('hybrid.png',15,'图5  同一输出后端下的混合控制流程与后续抓放接口')
h('切换条件与规划并发')
p('hybrid_enabled默认关闭。启用后，TCP到期望TCP的距离≤0.10 m进入PBVS，>0.16 m退出近距状态，迟滞避免边界反复切换。远距路径规划到距期望位0.06 m的接近点，预算3 s，执行超时30 s。规划回调按10 Hz检查待处理请求，服务调用期间不持有控制锁，其余回调仍能处理停止和反馈。')
h('对齐与目标丢失')
p('混合路径完成后先稳定0.3 s，再进入近距控制；目标位移超过0.04 m会使旧路径失效。当前默认位置死区为4 mm、姿态死区为0.02 rad，持续满足对齐条件0.35 s后进入ALIGNED并保持反馈位置；误差超过释放阈值后恢复跟踪。目标超时阈值为0.20 s，眼在手外默认采用stop策略；混合模式目标丢失时清除旧路径并保持。上述阈值是控制参数，不能视为实测终端抓取精度。')
p('固定相机默认target_offset为[0, −0.08, 0.10] m，包含避遮挡横向偏移和10 cm预抓取高度。因此ALIGNED表示到达设置的对准位，并不表示夹爪已经接触物体；最终接近动作应单独组织。')

page('七 工程难点与实现取舍')
table(['问题','处理方式','需要解释的边界'],[
 ('深度噪声与推理滞后','保留采集时间；有界深度缓存；ROI中值','拒绝无效结果，避免用旧位姿伪装新观测'),
 ('图像角度与夹爪方向不同','旋转框中线投影到工作平面','依赖水平面和物体几何先验'),
 ('指令生产与SDK通信不同步','生产消费解耦；按控制柜余量批量供点','队列容量和轮询Hz不是闭环带宽'),
 ('近目标抖动与逆解放大','连续软死区；DLS；关节限速限加速度','不能替代规划场景中的在线避碰'),
 ('目标丢失与重新捕获','超时状态机；保持；按模式选择恢复策略','眼在手外默认不执行搜索运动'),
 ('规划与伺服争用控制柜','混合模式共用队列和唯一输出后端','外部RobotHW不能同时独占SDK'),
],[3.8,6.1,7.5])
h('可展示的工程贡献')
p('感知侧：围绕旋转目标检测构建统一ROS消息，连接RGB-D时间配对、相机标定和抓取几何，把像素中心与方向转换成基座系抓取参数。对深度异常、无效方向和未标定宽度设置明确的拒绝条件。')
p('驱动侧：把厂商SDK封装为ros_control标准硬件接口，贯通MoveIt、关节轨迹控制器和真实控制柜；通过独立状态读取与供点线程，降低计算与网络通信之间的相互阻塞。')
p('控制侧：实现统一PBVS控制核心，使用雅可比阻尼伪逆、连续软死区、关节约束和有界位置队列；通过Gazebo与SDK两个输出后端复用算法，并加入目标超时与对齐保持状态。')
h('展示时应保留的限制')
p('ROS定时器运行在普通调度环境中，100 Hz或250 Hz属于目标节拍，不是硬实时保证。局部伺服没有直接复用MoveIt规划场景的碰撞检查，伺服工作区域及接近距离需先完成验证。')
p('Gazebo中使用的抓取附着插件是物理仿真辅助，不是实体夹爪的抓取检测。仿真成功率与实机成功率必须分别统计。仓库中的眼在手上、眼在手外以及table、depth模式也应按实际演示条件选择，避免展示图和启动参数不一致。')

page('八 实验口径与招聘展示讲述')
h('指标如何形成可复核的结果')
table(['指标','展示数值','随实验记录保存的内容'],[
 ('mAP@0.5','77.2%','数据集划分、类别、标注方式、模型版本与逐类AP'),
 ('位置与姿态误差','5.8 mm / 3.3°','参考真值、样本量、误差定义、均值或RMSE；姿态角含义'),
 ('模型预测耗时','平均5.4 ms','CPU/GPU、输入尺寸、预热、样本数、是否含前后处理'),
 ('单次运动规划','<100 ms','规划场景、起终点、规划器、失败次数及耗时分布'),
 ('综合抓取成功率','92.6%','成功次数/总次数、物体组成、成功判据及实机/仿真条件'),
],[3.5,3.2,10.7])
p('这些结果描述项目测试表现。招聘展示中应同时提供对应实验配置与原始记录，不将模型推理时间扩展为端到端延迟，也不将伺服死区或速度上限表述为实测精度。')
h('建议的演示顺序')
p('先用真实检测画面解释OBB为何适合任意平面摆放方向，再在RViz展示基座系抓取位姿、碰撞场景和规划轨迹。随后展示独立视觉伺服对目标位置变化的在线跟踪与目标丢失保持。混合模式可继续展示远距APPROACH到近距TRACKING的状态切换，结合运行日志说明阈值与目标丢失行为。')
h('招聘交流用项目介绍')
p('我们开发了AUBO i5智能分拣与视觉伺服系统，面向目标位置和摆放角度不确定的桌面分拣任务。系统将YOLO-OBB旋转检测与RealSense RGB-D数据结合，通过标定和坐标变换生成抓取位置与夹爪方向；使用MoveIt组织抓取、搬运和放置，并基于厂商SDK完成ros_control硬件接口和轨迹执行链路。')
p('在视觉闭环部分，我们根据目标与TCP的位姿误差，通过雅可比阻尼伪逆求关节速度，完成限速、限加速度与位置积分，再利用有界队列向执行后端供点。控制计算默认100 Hz，位置点按200 Hz生成，真机关节反馈50 Hz。混合模式通过MoveIt规划服务与同一执行队列，衔接长距离规划和近距离视觉闭环。')
p('项目实验中，检测mAP@0.5为77.2%，位置和姿态估计误差为5.8 mm和3.3°，平均模型预测耗时5.4 ms，单次规划耗时小于100 ms，综合抓取成功率92.6%。讲述时结合具体测试设备和样本口径，重点解释从检测结果到稳定执行的工程链路。')

page('九 源码索引与参数复核入口')
p('以下路径均相对于仓库根目录。正文中的S编号对应这些实现入口；频率与参数以当前启动时的实际覆盖值为准。')
table(['编号','源码或配置','核对内容'],[
 ('S1','aubo/aubo_perception/scripts/ultralytics_yolo_node.py\naubo/aubo_perception/config/ultralytics_yolo.yaml','OBB输出、原时间戳、非阻塞推理锁、10 Hz限频'),
 ('S2','aubo/aubo_perception/scripts/yolo_rgbd_target_node.py\naubo/aubo_perception/src/aubo_perception/grasp_geometry.py\naubo/aubo_perception/config/grasp_pipeline.yaml','深度缓存与匹配、反投影、顶面与中心、旋转框方向'),
 ('S3','aubo/aubo_sorting_core/src/color_sorting_task.cpp\naubo/aubo_sorting_core/src/target_tracking.cpp\naubo/aubo_sorting_core/src/pick_place.cpp','任务工作线程、目标缓存、抓放顺序、10 Hz等待'),
 ('S4','aubo/aubo_sorting_core/src/motion_executor.cpp\naubo/aubo_moveit_config/config/ompl_planning.yaml\naubo/aubo_sorting/config/yolo_sorting.yaml','RRTConnect、笛卡尔路径、时间参数化、规划预算'),
 ('S5','aubo/aubo_ros_control/src/aubo_hw_main.cpp\naubo/aubo_ros_control/src/aubo_hardware_interface.cpp\naubo/aubo_ros_control/include/aubo_ros_control/aubo_hardware_interface.h','250 Hz主循环、50 Hz状态读取、SDK供点、5 ms限幅常量'),
 ('S6','aubo/aubo_ros_control/src/visual_servo.cpp\naubo/aubo_ros_control/src/visual_servo_common.cpp\naubo/aubo_ros_control/src/direct_sdk_backend.cpp','PBVS、状态机、子步积分、队列、独立SDK输出线程'),
 ('S7','aubo/aubo_ros_control/config/visual_servo_common.yaml\naubo/aubo_ros_control/config/visual_servo_eye_to_hand.yaml\naubo/aubo_ros_control/config/controllers.yaml','100/200 Hz、死区、超时、姿态目标、状态发布频率'),
 ('S9','aubo/aubo_ros_control/src/hybrid_control.cpp\naubo/aubo_ros_control/config/hybrid_control.yaml','规划服务、轨迹校验、距离迟滞、短队列与状态切换'),
 ('S8','realsense-ros-development/realsense2_camera/launch/rs_camera.launch\naubo/aubo_perception/models/workspace_camera/model.sdf','真实相机默认30 Hz、固定仿真相机20 Hz'),
],[1,11.2,5.2])
h('相机与任务启动参数')
p('眼在手外应显式选择camera_mount:=eye_to_hand，并核对相机命名空间及标定TF；使用深度估计高度时选择height_mode:=depth。YOLO分拣便捷入口默认采用eye_in_hand和table，不能直接当作眼在手外RGB-D实验配置。')
p('开源仓库  https://github.com/2799063570/ros-mobile-manipulation-lab')
doc.core_properties.title='AUBO i5机械臂智能分拣与视觉伺服系统'
doc.core_properties.subject='招聘展示技术文档'
doc.core_properties.author='项目团队'
doc.save(OUT)
print(OUT)
